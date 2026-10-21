from pathlib import Path
from hashlib import sha256
from docx import Document
import json
import re

ROOT = Path('/home/fabian/Escritorio/8vo_Semestre/Proyecto_de_Titulo')
DOC = ROOT / 'Entregables_GeoCattle_2026-10-04/Informe_2_GeoCattle_con_Viabilidad_y_Flujo_Caja_2026-10-04.docx'
V3 = ROOT / 'geocattle-saas/docs/especificacion-v3'
PLAN = ROOT / 'geocattle-saas/docs/planificacion'
d = Document(DOC)
parts = []
def p(n): return d.paragraphs[n - 1].text
def section(n, title, source, operation, text):
    parts.append(f'## A{n:02d}. {title}\n\n**Lugar exacto:** {source}\n\n**Operación:** {operation}\n\n**Texto exacto para incorporar:**\n\n{text.strip()}\n')

parts.append(f'''# Actualización del informe GeoCattle

Fecha de corte: **21 de octubre de 2026**.

Este archivo contiene ubicaciones verificadas y texto listo para reemplazar o agregar en el informe. El DOCX original se conserva sin modificaciones.

**Archivo fuente principal:** [{DOC.name}]({DOC}). Es la versión más reciente del informe de desarrollo encontrada en el workspace, con actualización del 4 de octubre de 2026. No aplicar estos números de párrafo a otra copia del informe.

**Regla de localización:** P significa párrafo de texto contado desde 1 en el cuerpo del DOCX mediante `Document.paragraphs`; excluye celdas de tablas. T significa tabla contada desde 1 mediante `Document.tables`. Se proporciona además el título y un texto original para localizar el lugar desde Word o LibreOffice. No se asignan números de página, porque cambian al incorporar contenido.

**Fuentes de alcance:** `geocattle-saas/docs/especificacion-v3/00_ANALISIS_Y_ALCANCE.md`, `01_CASOS_DE_USO.md`, `02_REQUERIMIENTOS.md`, `03_MODELO_CONCEPTUAL.md`, `04_MODELO_LOGICO.md` y `05_MODELO_FISICO.md`. La decisión confirmada es que el veterinario realiza cambios directos. La especificación tiene 29 casos de uso, 44 requerimientos funcionales y 22 no funcionales; su modelo lógico objetivo tiene 48 tablas.

**Distinción necesaria:** modelo objetivo, código incorporado, prueba local y aceptación de campo son evidencias distintas. Los textos de alcance y diseño que siguen no certifican cumplimiento de los 66 requerimientos. La sección A06 consigna la implementación efectivamente verificada al cierre de esta actualización.

**Cambios de numeración:** los CU y RF del informe antiguo no equivalen por número a V3. Reemplazar los catálogos completos indicados abajo. En referencias históricas se escribirá `RF legado 01` o `CU legado 01`. En particular, las pruebas de radio `RF01`, `RF02` y `RF03` de T34 se renombran `RAD-01`, `RAD-02` y `RAD-03`, para diferenciarlas de los requerimientos funcionales `RF-01`, `RF-02` y `RF-03`.

| Bloque | Sección fuente | Cambio |
|---|---|---|
| A01–A05 | Resumen, I Introducción y alcance | Objetivos, cuatro actores y permisos. |
| A06 | Estado de desarrollo | Evidencia real y limitaciones. |
| A07–A09 | III Procesos y arquitectura | Aprobación/confianza, 29 CU y API. |
| A10–A12 | Modelo de datos y anexo 2 | Conceptual, lógico y físico incremental. |
| A13–A14 | IV Indicadores y V Plan de pruebas | KPIs por perfil y aceptación V3. |
| A15 | VII Cronograma | Gantt con fechas, dependencias y dos responsables. |
| A16 | Anexo 4 | Fichas completas de CU-01 a CU-29. |
| A17–A18 | Anexo 5 | 44 RF, 22 RNF y transición desde identificadores anteriores. |
| A19 | Índice y referencias a figuras | Rótulos y referencias consistentes. |
''')

section(1,'Resumen ejecutivo',f'Sección “Resumen ejecutivo”, después de P6, que empieza “{p(6)[:114]}…”','Agregar un párrafo; conservar P7–P9 y sus importes, porque esta actualización no recalcula el flujo económico.', '''La aplicación web incorpora cuatro perfiles: administrador de plataforma, dueño de cuenta cliente, trabajador y veterinario. El administrador consulta indicadores y auditoría global; el dueño administra varios predios dentro de su cupo, sus colaboradores y sus decisiones operativas. El trabajador registra propuestas sujetas a aprobación, salvo autorización de confianza para una acción específica. El veterinario modifica directamente la ficha técnica y sanitaria de los animales de los predios donde mantiene acceso vigente. Cada operación conserva autoría, contexto y trazabilidad. El alcance se documenta mediante 29 casos de uso, 44 requerimientos funcionales y 22 no funcionales, con implementación y aceptación verificadas por separado.''')
section(2,'Objetivo general',f'Sección “I Introducción y alcance” → “Objetivo general”, P33. Texto actual: “{p(33)}”','Reemplazar P33.', '''Diseñar e implementar GeoCattle como un sistema IoT y una aplicación web de monitoreo y gestión bovina con acceso diferenciado para administrador, dueño, trabajador y veterinario, que conserve ubicación, temperatura superficial, actividad, estado técnico, historial sanitario y responsabilidades por predio, y que apoye la supervisión mediante mapas, geocercas, alertas y aprobación trazable de cambios.''')
section(3,'Objetivos específicos',f'Sección “Objetivos específicos”, P35 a P39. Primera viñeta actual: “{p(35)}”','Reemplazar las cinco viñetas por las siguientes.', '''- Integrar el prototipo académico con el material disponible y especificar el collar del cliente con su BOM final y pruebas de banco y campo separadas.
- Implementar autenticación individual, segundo factor para administración, verificación reforzada de operaciones sensibles y autorización por cuenta, predio y función.
- Administrar predios, sectores, geocercas versionadas, bovinos identificados por DIIO, collares y gateways conservando asignaciones históricas.
- Diferenciar las propuestas del trabajador de los cambios efectivos, habilitar confianza granular y permitir registros veterinarios directos con historial sanitario inmutable.
- Persistir y presentar telemetría, eventos, medidas e historia sanitaria con calidad, unidad, fecha, autor y procedencia explícitas.
- Validar los cuatro paneles, la integridad de la migración, el aislamiento de datos, la restauración y el comportamiento sin Internet mediante pruebas reproducibles.
- Actualizar cronograma, documentación y criterios de aceptación, y demostrar la cadena con un nodo real antes de aceptar una preserie y de comprometer el lote de 70 collares.''')
section(4,'Alcance funcional',f'Sección “Alcance funcional”, P41 y P42; P41 empieza “{p(41)[:120]}…”','Reemplazar P41 y P42 por los siguientes párrafos.', '''Se incluyen la administración de clientes y recursos de plataforma, cuentas individuales, membresías por predio, cupos de predios, sectores según uso, geocercas de cuatro o más vértices distintos y registro bovino con DIIO, raza, color, sexo y nacimiento conocido o estimado. La edad se calcula a partir de nacimiento; peso y altura se conservan como mediciones fechadas. El collar físico, el gateway y el módulo funcional de software se administran y contabilizan como conceptos diferentes.

El dueño registra sus predios y colaboradores y decide las solicitudes de cambio del trabajador. Las operaciones sensibles de gestión de colaboradores requieren verificación reforzada. El trabajador puede reportar un hecho sanitario administrado, estado o reporte del bovino y cambio de batería. La confianza se concede por trabajador, predio y tipo de acción, y únicamente elimina la aprobación para la acción autorizada. El veterinario vigente registra y corrige directamente ficha técnica, vacunas, antibióticos y patologías, sin aprobación posterior del dueño.

La ficha sanitaria conserva originales, correcciones y anulaciones mediante registros vinculados. Eliminar acceso revoca una membresía y no borra una identidad compartida con otro predio ni su autoría histórica. Una solicitud pendiente se guarda en la base de datos, pero no modifica el registro efectivo hasta su aceptación. El registro de una administración sanitaria real pendiente mantiene visible su condición provisional y la advertencia de resguardo aplicable.

Se preservan asignaciones históricas, mapas, telemetría, alertas, mantenimiento y notificaciones. La temperatura es superficial y la actividad es un índice de movimiento; estos datos apoyan revisión presencial y no constituyen diagnóstico. El collar GeoCattle complementa el DIIO y no sustituye la identificación ni el registro oficial SAG. FMA permanece como extensión P2 y un documento interno no equivale a una emisión o confirmación oficial en SIPECweb.''')
scope=(V3/'00_ANALISIS_Y_ALCANCE.md').read_text()
matrix=scope.split('## Matriz de autorización\n\n',1)[1].split('\n## Prioridad',1)[0]
actors=scope.split('## Interpretación de los cuatro perfiles\n\n',1)[1].split('\n## Decisiones adoptadas',1)[0]
section(5,'Actores y reglas de autorización',f'Después del nuevo bloque “Alcance funcional”, antes del título “Presentación académica y despliegue al cliente” (P44 original).','Agregar un subtítulo “Perfiles de usuario y autorización” y el contenido siguiente.', '### Perfiles de usuario y autorización\n\n'+actors+'\n\n'+matrix+'\n\nLa autorización se valida en el servidor para cada petición. Ocultar una opción en el panel no reemplaza esa validación. Los colaboradores pueden participar en varios predios; el cambio de contexto conserva su identidad y exige membresía vigente.')

evidence=PLAN/'EVIDENCIA_IMPLEMENTACION.md'
evidence_text=evidence.read_text() if evidence.exists() else '''A 21 de octubre de 2026 existe una especificación V3 con cuatro perfiles, diagramas y modelos de datos. La ampliación de código, sus migraciones y los paneles se encuentran en integración y deben consignarse junto a sus resultados de ejecución antes del cierre del informe. La evidencia de compilación, pruebas SQL y pruebas automatizadas debe distinguirse de las comprobaciones aún no realizadas con usuarios, dispositivos y proveedores externos.

La aceptación de MFA, aislamiento, integridad y aprobación exige pruebas específicas. El ensayo local no acredita por sí solo rendimiento con 70 collares, disponibilidad durante siete días, entrega de correo externo, cobertura LoRaWAN, autonomía, resguardo clínico ni aceptación del productor. Los datos simulados se identifican y no sustituyen mediciones físicas.'''
section(6,'Estado real de desarrollo',f'Sección “Estado de desarrollo”, P50 y P51. P50 actual empieza “{p(50)[:140]}…”','Reemplazar P50 y P51 por este bloque, actualizado con la evidencia de integración disponible.',evidence_text)
section(7,'Proceso de aprobación y atención veterinaria',f'Sección “III Procesos de negocio y supervisión”, después de P243 y antes de “Casos de uso de la solución” (P245).','Agregar el subtítulo y proceso siguiente. Conservar los tres BPMN existentes de supervisión/alerta/configuración, actualizando sus actores cuando corresponda.', '''### P04 Registrar y resolver cambios sobre el bovino

El colaborador selecciona un predio autorizado y un animal de ese contexto. El sistema verifica identidad, membresía, acción permitida y versión base. Para un trabajador sin confianza vigente se persiste una solicitud pendiente con autor, fecha y contenido, se informa al dueño y se conserva sin efecto en el registro confirmado. El dueño revisa y aprueba o rechaza; el rechazo exige motivo. La aprobación comprueba nuevamente permisos y versión base y aplica el cambio una sola vez dentro de la misma transacción que guarda decisión y auditoría. Un conflicto conserva la solicitud para reformulación o rechazo.

Cuando existe confianza específica, el sistema valida su vigencia al aplicar y registra qué concesión permitió el cambio directo. El veterinario realiza cambios técnicos y sanitarios directos en sus predios autorizados; su operación conserva confirmación, autoría y auditoría y no pasa por aprobación del dueño. La planificación de una vacuna no demuestra que haya sido administrada. Las correcciones y anulaciones sanitarias agregan registros relacionados y conservan el original.

La confirmación de una decisión de negocio y la doble verificación de seguridad son controles diferentes. Una ventana de verificación válida no concede aprobación sobre una solicitud ni habilita funciones ajenas al perfil.''')
cu=(V3/'01_CASOS_DE_USO.md').read_text()
catalog=cu.split('## Catálogo\n\n',1)[1].split('\n## CU-01',1)[0] if '## Catálogo\n\n' in cu else cu[cu.index('| ID |'):cu.index('\n## CU-01')]
section(8,'Catálogo y diagramas de casos de uso',f'Sección “Casos de uso de la solución” (P245). Sustituir T27, encabezado “Caso / Actor y resultado esperado”, Figura 7 ubicada entre P245 y P247 y comentario P249.','Sustituir la tabla y el comentario por el siguiente catálogo. Reemplazar Figura 7 por la vista general e insertar vistas por perfil como Figuras 7a–7g, con los rótulos siguientes.', '''Los casos se identifican como CU-01 a CU-29 de la especificación V3. Los cuatro actores humanos son administrador, dueño, trabajador y veterinario. Los servicios de ingestión y avisos son actores técnicos. Autenticación válida y autorización por contexto son precondiciones; no se dibuja inicio de sesión como inclusión de todas las funciones. Las asociaciones UML no representan un orden de ejecución.

'''+catalog+'''

| Figura nueva | Rótulo exacto | Archivo para insertar |
|---|---|---|
| 7 | Casos de uso V3: vista general de actores y funciones. Elaboración propia. | `diagramas-casos-uso/01_general.svg` |
| 7a | Casos de uso del administrador de plataforma. Elaboración propia. | `diagramas-casos-uso/02_administrador.svg` |
| 7b | Casos de uso del dueño: predios, colaboradores y decisiones. Elaboración propia. | `diagramas-casos-uso/03_dueno_gestion.svg` |
| 7c | Casos de uso del dueño: monitoreo y salud. Elaboración propia. | `diagramas-casos-uso/04_dueno_monitoreo_salud.svg` |
| 7d | Casos de uso del trabajador: consultas y propuestas. Elaboración propia. | `diagramas-casos-uso/05_trabajador.svg` |
| 7e | Casos de uso del veterinario: modificación directa de fichas. Elaboración propia. | `diagramas-casos-uso/06_veterinario.svg` |
| 7f | Casos de uso de aprobación y confianza por acción. Elaboración propia. | `diagramas-casos-uso/07_aprobacion.svg` |
| 7g | Casos de uso de servicios técnicos de telemetría, alertas y avisos. Elaboración propia. | `diagramas-casos-uso/08_actores_tecnicos.svg` |

La carpeta de estas figuras es `geocattle-saas/docs/especificacion-v3/diagramas-casos-uso/`. El archivo multipágina editable es `CASOS_DE_USO_V3.drawio`. La entrega de avisos externos y la extensión FMA conservan sus dependencias de proveedor y documentación oficial; una figura no constituye evidencia de ejecución.''')
section(9,'Arquitectura y contrato de aplicación',f'Sección “Componentes lógicos y arquitectura de software”, después de P255; sección “Contratos de la API”, P258.','Agregar el primer párrafo bajo arquitectura y sustituir P258 por los restantes.', '''La aplicación organiza identidad y seguridad, administración de plataforma, gestión predial, ganado, IoT, salud, aprobación y auditoría como dominios de un backend modular. Los cuatro paneles consumen la API mediante la identidad autenticada y el contexto autorizado. El panel global usa agregaciones administrativas; los paneles prediales acceden únicamente a cuentas y predios vigentes. No se requiere un microservicio por tabla ni una base separada por perfil.

Toda escritura valida identidad, ámbito, permiso y contenido en el servidor. Las operaciones sensibles verifican además el segundo factor o desafío correspondiente. Una solicitud de cambio conserva estado, versión base y decisión, y solo una aplicación confirmada modifica los datos efectivos. La aprobación no acepta contenido distinto del revisado. Las respuestas de error no presentan operaciones como exitosas.

El esquema físico se actualiza mediante migraciones versionadas y aditivas sobre la aplicación existente, conservando datos e identificadores heredados. El esquema objetivo V3 sirve como referencia de diseño y no se aplica como un reemplazo destructivo del esquema operativo. La simulación mantiene origen y entorno diferenciados; las pruebas negativas incluyen acceso a identificadores de otra cuenta y de otro predio.''')
concept=(V3/'03_MODELO_CONCEPTUAL.md').read_text()
concept_body=concept[concept.index('```mermaid'):]
section(10,'Modelo conceptual entidad-relación',f'Sección “Modelo lógico de datos” (P260), antes de Figura 9.','Agregar un subtítulo “Modelo conceptual entidad-relación” y el siguiente contenido. La imagen resultante se incorpora como Figura 9a; usar la página conceptual del archivo MODELO_V3.drawio para editar.', '''El modelo conceptual separa usuario, acceso de plataforma, cuenta de dueño, membresía predial, animal y permanencia, collar, geocerca, solicitud, decisión, salud y auditoría. La relación usuario–predio es muchos a muchos mediante membresía; la relación animal–predio es histórica mediante permanencia. Las relaciones temporales no permiten atribuir una lectura antigua al animal que utiliza el collar en el presente.

'''+concept_body)
section(11,'Modelo lógico resumido',f'Sección “Modelo lógico de datos”, Figura 9 y P263; P263 actual empieza “{p(263)[:130]}…”','Reemplazar Figura 9 por la página lógica de MODELO_V3.drawio o su fuente MODELO_LOGICO_V3.mmd. Reemplazar P263 por los párrafos siguientes. Conservar P265–P267 y añadir las reglas finales.', '''El modelo lógico objetivo V3 contiene 48 tablas agrupadas en identidad, predios, ganado y espacio, IoT, operación, aprobación, salud, administración y extensión FMA. El usuario es una identidad única; la administración global se representa aparte de la cuenta cliente. Una cuenta de dueño posee varios predios dentro de su cupo. Trabajadores y veterinarios obtienen acceso por membresía y pueden participar en más de un predio.

El animal mantiene DIIO y descripción; su presencia por predio se conserva en intervalos. Las asignaciones de collar y pastoreo conservan su vigencia. La telemetría distingue observación y recepción, sesión y secuencia, y resuelve el contexto de captura. Una lectura puede tener varias recepciones de gateway sin convertirse en varias lecturas. La historia sanitaria y sus tratamientos, medidas, reportes y agenda se separan de las observaciones del sensor.

Las solicitudes, decisiones y aplicaciones registran el flujo del trabajador y la evidencia del cambio directo veterinario. La auditoría de seguridad y negocio se conserva aparte de los eventos técnicos y de las métricas de peticiones. Los módulos funcionales habilitados y los dispositivos físicos se contabilizan por separado. El diccionario lógico completo se presenta en el anexo 2; el esquema físico desplegado se identifica por migración y no por el número de tablas del diseño objetivo.

Una geocerca tiene al menos cuatro vértices distintos y cierre automático. El borde pertenece al interior. Peso y altura se guardan con fecha, autor y unidad; la edad es derivada. Ninguna cardinalidad reemplaza la validación de ámbito, el control de intervalos o la autorización en API. Un evento sanitario confirmado se corrige o anula por sucesor y conserva el original.''')
logical=(V3/'04_MODELO_LOGICO.md').read_text()
logical_names=re.findall(r'^### `([^`]+)`',logical,re.M)
physical={}
for sql in sorted((ROOT/'geocattle-saas/database/migrations').glob('*.sql')):
    for name in re.findall(r'CREATE TABLE (?:IF NOT EXISTS )?([a-z_]+)',sql.read_text(),re.I):
        physical[name]=sql.name
correspondence='| Tabla del modelo objetivo | Tabla física operativa | Migración de creación |\n|---|---|---|\n'+'\n'.join(f"| `{name}` | `{'auth_challenges' if name=='security_challenges' else name}` | `{physical.get('auth_challenges' if name=='security_challenges' else name,'REVISAR')}` |" for name in logical_names)
assert 'REVISAR' not in correspondence
(PLAN/'CORRESPONDENCIA_MODELO_V3.md').write_text('# Correspondencia del modelo lógico V3 con las migraciones operativas\n\nInventario obtenido de las sentencias CREATE TABLE de las migraciones 001, 002 y 003. Las tablas preexistentes pueden recibir columnas, restricciones y triggers adicionales en 003. Esta tabla compara nombres y procedencia; no equipara todas las columnas de la propuesta con las del esquema operativo.\n\n'+correspondence+'\n\n`security_challenges` de la propuesta se implementa como `auth_challenges`. El historial de cambios operativo se consulta en `003_application_v3.sql`.\n')
section(12,'Diccionario lógico y modelo físico',f'Anexo 2: título P426 “{p(426)}”, P428–P435, T44 “Entidad / Campos y tipos propuestos / Restricciones” y T45 del mismo encabezado bajo P431.','Sustituir las dos tablas resumidas y sus párrafos introductorios por el siguiente texto. Incorporar como subanexo el diccionario completo transcrito después. Conservar política de retención como propuesta y no habilitar borrado histórico automático.', '''El anexo distingue el modelo lógico objetivo V3 y el esquema físico operativo. El diccionario de 48 tablas es una especificación de diseño. La aplicación existente se actualiza mediante una migración aditiva que conserva el esquema operativo y los identificadores. La validación registra por separado creación de base vacía, actualización de una copia con datos heredados e idempotencia de migración. No se ejecuta un segundo esquema como sustitución automática de los datos del sistema.

Las restricciones comprenden claves foráneas y de unicidad, intervalos no superpuestos, coherencia de predio, geometría válida y protección de historia. La autorización de usuario, el segundo factor, la membresía y la confianza se validan además en la API. Los índices espaciales y temporales se verifican con el volumen de prueba. TimescaleDB se mantiene como alternativa de series temporales; no se declara instalada una hypertable sin migración y evidencia específicas.

El diccionario de diseño se conserva en `geocattle-saas/docs/especificacion-v3/04_MODELO_LOGICO.md`; el DDL objetivo está en `MODELO_FISICO_V3.sql`. El inventario exacto de tablas, columnas y restricciones efectivamente incorporadas en esta actualización consta en la migración operativa `database/migrations/003_application_v3.sql` y en su registro de verificación. Los nombres del DDL objetivo y de la migración se comparan antes de trasladarlos al informe.

### Correspondencia de nombres y procedencia física

Las migraciones 001, 002 y 003 definen respectivamente 18, 1 y 29 tablas; el inventario operativo suma 48. La tabla `security_challenges` de la propuesta corresponde a `auth_challenges` en la implementación. Las tablas conservadas reciben ampliaciones mediante ALTER TABLE; conservar su nombre no implica que sus atributos permanezcan iguales.

'''+correspondence+'''

### Diccionario lógico objetivo V3

'''+logical.split('\n',1)[1])
section(13,'KPIs de administración y de operación',f'Sección “IV Indicadores KPI y medición SMART”, después de P289 y antes de T30 de encabezados “KPI / Cálculo y meta / Ventana y evidencia”.','Agregar el siguiente bloque y tabla. Conservar K01–K12 y las mediciones físicas existentes como indicadores de validación del prototipo.', '''Los indicadores de plataforma se consultan por período y por cuenta de dueño. Se distingue stock actual de actividad diaria y se muestra la fecha de corte. El dueño consulta el contexto propio; el administrador consulta agregaciones globales. No se duplican usuarios por pertenecer a varios predios ni alertas por reintentos de aviso.

| ID | Indicador | Definición y filtro |
|---|---|---|
| KP-V3-01 | Usuarios únicos | Número de identidades activas distintas; la vista por dueño considera la cuenta y sus membresías vigentes sin duplicación. |
| KP-V3-02 | Bovinos registrados | Animales con pertenencia/permanencia activa, filtrados por cuenta y predio. Informar altas, bajas y fecha de corte. |
| KP-V3-03 | Predios y sectores | Contar predios y sectores activos por separado. No sumar magnitudes de entidades diferentes. |
| KP-V3-04 | Collares | Distinguir inventario, asignados/instalados y retirados; un collar sin animal no cuenta como instalado. |
| KP-V3-05 | Módulos de software | Funcionalidades habilitadas por predio; no equivalen al número de collares. |
| KP-V3-06 | Alertas por día | Ocurrencias según fecha de apertura en zona del predio. Intentos de correo no aumentan alertas. |
| KP-V3-07 | Consumo de API | Peticiones, errores, bytes y p95 de duración por período y cuenta, sin almacenar secretos ni cuerpos de petición. |
| KP-V3-08 | Solicitudes pendientes | Cambios pendientes por predio, autor, tipo y antigüedad; conservar distinto el conteo de aplicados y rechazados. |

CPU, memoria y disco proceden del sistema de observabilidad de la instancia y no de una inferencia desde el número de usuarios. Las instalaciones desconectadas muestran su última información disponible. Las metas de latencia, disponibilidad y usabilidad siguen pendientes hasta medirlas en los escenarios de aceptación.''')
section(14,'Pruebas V3 y criterios de salida',f'Sección “V Plan de pruebas de software e integración”, después de P304 y antes de “Ambientes y criterio de salida” (P305). T34 radio conserva sus procedimientos y cambia etiquetas RF01/RF02/RF03 por RAD-01/RAD-02/RAD-03.','Agregar la tabla siguiente; conservar las pruebas previas de sensores, radio, montaje y energía. Agregar al final de P306 el criterio de salida final.', '''| ID | Escenario | Resultado exigido |
|---|---|---|
| V3-S01 | Administrador con/sin segundo factor, código vencido y repetido | Sin segundo factor verificado no existe acceso administrativo utilizable. |
| V3-S02 | Dos dueños, predios distintos y colaborador compartido | API rechaza consultas y escrituras con IDs fuera de ámbito; cada contexto conserva su rol. |
| V3-S03 | Invitación vigente, vencida, repetida y correo distinto | Enrolamiento único y autorizado; no existe contraseña inicial compartida. |
| V3-S04 | Alta/cambio/baja de colaborador con reto incorrecto o reutilizado | Rechazo sin cambio; reto válido ligado a la acción y contenido exactos. |
| V3-S05 | Solicitud del trabajador pendiente, aprobada y rechazada | Pendiente no cambia ficha efectiva; aprobación aplica una vez y rechazo conserva motivo. |
| V3-S06 | Dos revisores o versión base obsoleta | No hay doble aplicación ni sobrescritura silenciosa; conflicto visible. |
| V3-S07 | Confianza revocada entre captura y envío | Se comprueba al aplicar y se requiere aprobación cuando ya no existe permiso. |
| V3-S08 | Veterinario registra/corrige/anula ficha y tratamiento | Cambio directo autorizado y auditado; originales sanitarios recuperables. |
| V3-S09 | Geocerca con tres vértices, repetidos y autointersección | Se rechaza geometría inválida y se conserva versión activa. |
| V3-S10 | Collar reasignado y paquete atrasado/duplicado | Historia atribuida según captura; sin nueva lectura ni evento por duplicado. |
| V3-S11 | KPIs por dueño y usuario en varios predios | Conteos con denominadores correctos y filtro de cuenta válido. |
| V3-S12 | Migración en base vacía y copia de V1 con datos | Historial e IDs conservados, restricciones verificadas y migración versionada. |
| V3-S13 | 70 nodos a T=300 s y cinco clientes durante 10 min | Medir RNF-08–RNF-10 sin extrapolar desde una prueba unitaria. |
| V3-S14 | Respaldo y restauración, corte WAN y reintentos | RPO/RTO medidos y consultas locales disponibles con colas recuperables. |
| V3-S15 | Correo y notificación con proveedor real | Entrega/fallo/reintento trazables; fallo externo no elimina solicitud o alerta. |
| V3-S16 | Tareas de usuario en móvil y escritorio | Cumplimiento medido de RNF-14, con resultado por tarea y perfil. |

El cierre académico exige evidencia de los flujos que se demuestran y ausencia de defectos críticos. El piloto requiere además seguridad, restauración y operación local verificadas. La aceptación de campo se mantiene separada y exige cobertura, autonomía, montaje, continuidad y revisión del productor. Una compilación correcta o una prueba SQL aislada no acreditan todas estas condiciones.''')

plan=json.loads((PLAN/'gantt_v3.json').read_text())
task_table='| ID | Actividad | Responsable propuesto | Inicio | Fin | Depende de |\n|---|---|---|---|---|---|\n'+'\n'.join(f"| {t['id']} | {t['task']} | {t['owner']} | {t['start']} | {t['end']} | {', '.join(t['deps']) or '—'} |" for t in plan['tasks'])
section(15,'Cronograma actualizado',f'Sección “VII Cronograma ajustado y organización”, P340, T38 de encabezados “Fase / Semanas / Dependencia y entrega / Responsable”, P343 y T39 “Hito / Evidencia necesaria”.','Reemplazar P340, T38, P343 y T39 por el bloque siguiente. Conservar P345 de identificación de ambos integrantes.', '''La línea base V3 comienza el 21 de octubre de 2026 y reserva doce semanas desde el corte para consolidación, verificación y preserie condicionada. La vista por semanas de lunes a domingo utiliza S1 parcial desde el 19 de octubre y S13 parcial hasta el 12 de enero de 2027. El software y la demostración se revisan hasta el 15 de diciembre; la documentación académica se consolida al 18 de diciembre. La campaña predial de cuatro semanas se propone del 16 de diciembre al 12 de enero, sujeta a equipos, autorización del predio y disponibilidad. No se dispone de un plazo académico oficial en las fuentes, por lo que estas fechas deben ratificarse con el calendario de la asignatura.

Fabián Del Villar coordina datos, API y seguridad; Boris Baier coordina interfaz, integración IoT y validación, con revisión cruzada. La planificación contiene 25 tareas y 410 horas estimadas, sin superar seis horas planificadas por persona y día. Los esfuerzos representan estimaciones y no horas efectivamente ejecutadas. Se usan días de lunes a viernes y se proponen como no laborables 8 y 25 de diciembre y 1 de enero; el equipo debe confirmar su disponibilidad de fin de año.

'''+task_table+'''

La ruta de aceptación de software comprende migración, identidad y permisos, flujos de aprobación y salud, KPIs, carga, restauración, avisos y regresión. La ruta física comprende cierre de diseño, unidades iniciales, interoperabilidad, pruebas de banco, campaña de cobertura/autonomía y aceptación de preserie. Una ruta no sustituye a la otra. La preserie no comienza hasta superar los hitos técnicos y de uso; un atraso de equipos o una prueba fallida desplaza las tareas dependientes.

| Hito | Fecha objetivo | Evidencia necesaria |
|---|---|---|
| H1 Base V3 incorporada localmente | 21/10/2026 | Migración, API y cuatro paneles con comprobaciones locales registradas; sin aceptación externa implícita. |
| H2 Versión académica verificada | 15/12/2026 | Seguridad, flujo de negocio, restauración, regresión y demostración con nodo real aprobados. |
| H3 Entrega documental académica | 18/12/2026 | Informe, Gantt, manual, inventario y matriz de resultados ejecutados/pendientes. |
| H4 Preserie aceptada | 12/01/2027, condicionada | Campaña de cuatro semanas, continuidad de siete días, cobertura, autonomía, montaje y acta del productor. |
| H5 Cliente con 70 collares | Sin fecha comprometida | H4 aprobado, costo actualizado, fabricación, instalación y capacitación. |

La carta editable se adjunta como `Carta_Gantt_GeoCattle_V3_2026-10-21.xlsx`, y su vista vectorial como `Carta_Gantt_GeoCattle_V3_2026-10-21.svg`, ambas en `geocattle-saas/docs/planificacion/`. Las fechas de fin, esfuerzos y barras se recalculan al cambiar entradas; las dependencias y la capacidad deben revisarse después de cualquier reprogramación. El FMA y el intercambio oficial SIPEC permanecen como extensión P2 sin fecha comprometida.''')
section(16,'Fichas completas de los casos de uso',f'Anexo 4, desde título P445 “{p(445)}” hasta P487 inclusive; excluir “Guion de presentación académica” P488–P489, que se conserva y se amplía por perfiles.','Reemplazar las fichas antiguas CU01–CU10 por el bloque V3 completo siguiente. El guion se conserva como prueba académica y deja de ser un caso de negocio numerado.',cu[cu.index('## CU-01'):])
req=(V3/'02_REQUERIMIENTOS.md').read_text()
rf=req.split('## Funcionales\n\n',1)[1].split('\n## No funcionales',1)[0]
rnf=req.split('## No funcionales\n\n',1)[1].split('\n## Trazabilidad',1)[0]
section(17,'Requerimientos funcionales V3',f'Anexo 5 “Trazabilidad de requerimientos funcionales” P491, T47 encabezada “RF / Capacidad / Casos / Pruebas o indicador” y comentario P493.','Reemplazar T47 y P493 por el siguiente texto y catálogo. No reasignar por coincidencia numérica los casos del informe antiguo.', '''Se adopta la numeración RF-01 a RF-44 de V3, vinculada a CU-01 a CU-29. Los criterios son condiciones verificables de aceptación, no resultados ya obtenidos. RF-43 corresponde a extensión P2. Se conservan las reglas heredadas de GNSS, intervalos, temperatura superficial y actividad, junto con las pruebas de hardware y radio de sus anexos.

'''+rf)
old_map=[('01','Registro individual','RF-14'),('02','Ubicación periódica','RF-22, RF-35'),('03','Zonas autorizadas','RF-13, RF-17'),('04','Detectar salida','RF-37'),('05','Alertar tras detección','RF-37, RF-40'),('06','Identificar animal y evento','RF-23, RF-37'),('07','Historial de ubicación','RF-38'),('08','Actividad','RF-23, RF-35'),('09','Baja actividad prolongada','RF-37'),('10','Temperatura asociada','RF-23, RF-35'),('11','Alerta térmica configurada','RF-37'),('12','Sensores en alerta','RF-23, RF-37'),('13','Propietario destinatario','RF-40'),('14','Consulta digital','RF-22'),('15','Historial de eventos','RF-38'),('16','Intervalo configurable','RF-39'),('17','Asociación DIIO-módulo','RF-16'),('18','Pérdida de comunicación','RF-37'),('19','Severidad de alertas','RF-37'),('20','Rebaño, sector y animal','RF-22, RF-38'),('21','Contexto del evento','RF-23, RF-37'),('22','Historial por período','RF-38'),('23','Orientación carne/leche','RF-14, RF-23 y atributo de orientación conservado'),('24','Recomendación preventiva','RF-37 y regla heredada no diagnóstica')]
map_table='| RF legado | Capacidad previa | Referencia V3 |\n|---|---|---|\n'+'\n'.join(f'| {a} | {b} | {c} |' for a,b,c in old_map)
section(18,'Requerimientos no funcionales y equivalencias',f'Subtítulo “Trazabilidad no funcional y cierre de decisiones” P495, T48 “RNF / Condición de diseño / Validación”. Conservar “Decisiones a cerrar antes de fabricar o instalar” P497 y T49.','Sustituir T48 por los 22 RNF siguientes y agregar la tabla de equivalencia funcional. Conservar los requisitos físicos heredados de cobertura, autonomía, intemperie, montaje y costo como controles HW/RAD/ME/EN, que V3 no sustituye.',rnf+'''

### Transición desde requerimientos del informe anterior

'''+map_table+'''

La equivalencia conserva la capacidad y no prueba cumplimiento. RF-14 y RF-23 deben mantener la orientación productiva carne/leche/doble propósito/otro ya definida en los datos heredados. Las pruebas RAD-01/RAD-02/RAD-03 siguen evaluando radio; no corresponden a RF-01/RF-02/RF-03 de autenticación. Las metas clínicas, eléctricas y de radio se aceptan con evidencia propia.''')
section(19,'Índice, figuras y anexos',f'“Contenido del informe”, P11–P20; rótulos P247 Figura 7 y P262 Figura 9; encabezados de anexos P426, P445 y P491.','Actualizar el índice al terminar la edición y usar los siguientes nombres de apartados y rótulos. No reutilizar las páginas antiguas del índice como localizadores.', '''- Perfiles de usuario y autorización.
- P04 Registrar y resolver cambios sobre el bovino.
- Casos de uso V3 de la solución.
- Modelo conceptual entidad-relación.
- Modelo lógico y correspondencia con el esquema físico.
- Indicadores de plataforma y de operación.
- Pruebas V3 y criterios de salida.
- VII Cronograma V3 y organización del trabajo.
- Anexo 2 Diccionario lógico V3 y esquema físico operativo.
- Anexo 4 Fichas CU-01 a CU-29.
- Anexo 5 Requerimientos RF-01 a RF-44 y RNF-01 a RNF-22.

Figura 9. Modelo lógico objetivo V3 de GeoCattle: entidades, claves y relaciones. Elaboración propia.

Figura 9a. Modelo conceptual entidad-relación de GeoCattle. Elaboración propia.

Tabla 9b. Correspondencia de nombres entre modelo objetivo V3 y migraciones del esquema operativo. Elaboración propia a partir del inventario SQL.

Figura 10a. Carta Gantt V3: tareas, dependencias y responsables desde el 21 de octubre de 2026. Elaboración propia.''')

parts.append(f'''## Registro de fuente y comprobación

- SHA-256 del DOCX fuente: `{sha256(DOC.read_bytes()).hexdigest()}`.
- Estructura inspeccionada: {len(d.paragraphs)} párrafos de texto y {len(d.tables)} tablas.
- T38 contiene la planificación anterior de ocho semanas; se preserva el archivo original.
- La totalidad de los 29 CU, los 44 RF, los 22 RNF y las 48 entradas del diccionario objetivo se incorpora o referencia en sus bloques específicos.
- La fecha final de campo es condicionada; no se modifica retrospectivamente el presupuesto ni la fecha de una evidencia anterior.
- Aplicar estas ediciones a una copia identificada del informe y actualizar índice/numeración al terminar. El archivo actual es una guía exacta de edición, no una modificación automática del DOCX.
''')
output=ROOT/'actualizacion_informe.md'
output.write_text('\n\n'.join(parts))
assert len(re.findall(r'^## A\d\d\.',output.read_text(),re.M))==19
print(json.dumps({'output':str(output),'bytes':output.stat().st_size,'blocks':19,'sourceParagraphs':len(d.paragraphs),'sourceTables':len(d.tables)},ensure_ascii=False))
